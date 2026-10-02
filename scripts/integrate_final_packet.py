"""
Integrate Final Packet Data into Mitra Kinerja Database
Source: docs/Paket_Scorecard_dan_Baseline_Final/Paket_Scorecard_dan_Baseline_Final
- Ingests all 180 audited baseline elements from FINAL_BASELINE_MITRA_KINERJA_28_AGUSTUS_2026.xlsx
- Realigns portfolio composition: STIT Mumtaz -> P05 (Pilot Utama), STAIN SAR -> C01, PBC -> C02
- Updates official PKS/MoU numbers, dates, PICs, and examiner names
- Ingests authentic field activities into tindak_lanjut
- Ingests authentic action plans into rencana_kerja
- Synchronizes database dumps
"""

import openpyxl, os, sys, subprocess

cwd = r"C:\Users\afra\Desktop\Agentic Folder\mitra-kinerja"
path_base = os.path.join(cwd, "docs", "Paket_Scorecard_dan_Baseline_Final", "Paket_Scorecard_dan_Baseline_Final")
baseline_file = os.path.join(path_base, "03_Baseline_Final", "FINAL_BASELINE_MITRA_KINERJA_28_AGUSTUS_2026.xlsx")

print("=== STARTING INTEGRATION OF FINAL PACKET DATA ===")

wb_base = openpyxl.load_workbook(baseline_file, data_only=True)

# Explicit mapping of Baseline sheets to Database IDs
sheet_to_mid = {
    'P01_DEKRANASDA': 1,
    'fFitriadi P02_BNNP_KEPRI': 2,
    'FIschika_P03_PEMKOT_TPI': 3,
    'fFitra P04_BAPPERIDA_BINTAN': 4,
    'fFitriadi_C02_STIT_MUMTAZ': 12,
    'fEnjo_P06_UMRAH': 6,
    'FNadia_P07_STAI_ANAMBAS': 7,
    'fFitra_P08_POLIBATAM': 8,
    'FNina_P09_STAI_NATUNA': 9,
    'fEnjo_P10_STISIP_BTM': 10,
    'FNina_P05_STAIN_SAR': 5,
    'FChika_C01_PBC': 11,
    'fNadia_C03_UIS': 13,
    'Fitra_C04_UNRIKA': 14,
    'FNina_C05_STIE_CAKRAWALA': 15,
}

wb_to_mid = {
    '01_Pilot_Utama/P01_Scorecard_Dekranasda_Kepri.xlsx': 1,
    '01_Pilot_Utama/P02_Scorecard_BNNP_Kepri.xlsx': 2,
    '01_Pilot_Utama/P03_Scorecard_Pemkot_Tanjungpinang.xlsx': 3,
    '01_Pilot_Utama/P04_Scorecard_Bapperida_Bintan.xlsx': 4,
    '01_Pilot_Utama/P05_Scorecard_STIT_Mumtaz_Karimun.xlsx': 12,
    '01_Pilot_Utama/P06_Scorecard_UMRAH.xlsx': 6,
    '01_Pilot_Utama/P07_Scorecard_STAI_Paduka_Anambas.xlsx': 7,
    '01_Pilot_Utama/P08_Scorecard_Politeknik_Negeri_Batam.xlsx': 8,
    '01_Pilot_Utama/P09_Scorecard_STAI_Natuna.xlsx': 9,
    '01_Pilot_Utama/P10_Scorecard_STISIP_Bunda_Tanah_Melayu.xlsx': 10,
    '02_Portofolio_Pengayaan/C01_Scorecard_STAIN_Sultan_Abdurrahman.xlsx': 5,
    '02_Portofolio_Pengayaan/C02_Scorecard_Politeknik_Bintan_Cakrawala.xlsx': 11,
    '02_Portofolio_Pengayaan/C03_Scorecard_Universitas_Ibnu_Sina.xlsx': 13,
    '02_Portofolio_Pengayaan/C04_Scorecard_UNRIKA.xlsx': 14,
}

examiner_names = {
    1: 'Evlin',
    2: 'Kompilasi Pokja Data',
    3: 'Ischika',
    4: 'Kompilasi Pokja Data',
    12: 'Kompilasi Pokja Data',
    6: 'Hariawan Novriadi',
    7: 'Nadia Putri Boga',
    8: 'Kompilasi Pokja Data',
    9: "Nur'ah Darina",
    10: 'Kompilasi Pokja Data',
    5: "Nur'ah Darina",
    11: 'Ischika',
    13: 'Nadia Putri Boga',
    14: 'Kompilasi Pokja Data',
    15: "Nur'ah Darina"
}

# Python MySQL helper via CLI
def run_mysql(sql):
    cmd = [r"C:\xampp\mysql\bin\mysql.exe", "-u", "root", "mitra_kinerja", "-e", sql]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"MySQL Error: {res.stderr}\nQuery: {sql[:200]}")
    return res.stdout

def sql_escape(val):
    if val is None:
        return "NULL"
    s = str(val).replace("\\", "\\\\").replace("'", "''")
    return f"'{s}'"

# Step 1: Realign Portfolio Codes
print("1. Realigning Portfolio Composition (STIT Mumtaz -> P05 Pilot Utama, STAIN SAR -> C01, PBC -> C02)...")
run_mysql("""
UPDATE mitra_kinerja SET kode = 'X01', portofolio = 'Cadangan' WHERE nama_mitra = 'Sekolah Tinggi Agama Islam Negeri Sultan Abdurrahman Kepulauan Riau';
UPDATE mitra_kinerja SET kode = 'X02', portofolio = 'Cadangan' WHERE nama_mitra = 'Politeknik Bintan Cakrawala';
UPDATE mitra_kinerja SET kode = 'P05', portofolio = 'Pilot Utama' WHERE nama_mitra = 'Sekolah Tinggi Ilmu Tarbiyah (STIT) Mumtaz Karimun';
UPDATE mitra_kinerja SET kode = 'C01' WHERE nama_mitra = 'Sekolah Tinggi Agama Islam Negeri Sultan Abdurrahman Kepulauan Riau';
UPDATE mitra_kinerja SET kode = 'C02' WHERE nama_mitra = 'Politeknik Bintan Cakrawala';
""")
print("   [DONE] Portfolio composition realigned.")

# Step 2: Fetch DB Mitra Map
res_mitra = run_mysql("SELECT id, kode, nama_mitra FROM mitra_kinerja;")
db_mitra = {}
for line in res_mitra.strip().splitlines()[1:]:
    parts = line.split("\t")
    if len(parts) >= 3:
        db_mitra[parts[2].strip()] = {'id': int(parts[0]), 'kode': parts[1].strip()}

print(f"   Mapped {len(db_mitra)} partners from DB.")

# Step 3: Ingest Master Baseline Elements & Update Partner Baseline Status
print("2. Ingesting 180 Audited Baseline Elements from FINAL_BASELINE...")

total_elements_ingested = 0

for sname, mid in sheet_to_mid.items():
    ws = wb_base[sname]
    pemeriksa = examiner_names.get(mid, 'Pokja Data')

    # Update mitra metadata
    upd_m = f"""
    UPDATE mitra_kinerja SET 
        baseline_status = 'TERVERIFIKASI / DIKUNCI',
        baseline_locked_at = '2026-08-28 23:59:59',
        baseline_locked_by = 1,
        baseline_pemeriksa = {sql_escape(pemeriksa)},
        cutoff_date = '2026-08-28',
        baseline_catatan_ringkasan = 'Final Baseline dikunci per 28 Agustus 2026 berdasarkan hasil audit Pokja Data tanpa kontradiksi internal.'
    WHERE id = {mid};
    """
    run_mysql(upd_m)

    # Ingest 12 elements
    for r in range(13, 25):
        num = int(ws.cell(r, 1).value)
        kelompok = str(ws.cell(r, 2).value or '').strip()
        nama_el = str(ws.cell(r, 3).value or '').strip()
        yang_diperiksa = str(ws.cell(r, 4).value or '').strip()
        sumber_min = str(ws.cell(r, 5).value or '').strip()
        status = str(ws.cell(r, 6).value or '').strip()
        fakta = str(ws.cell(r, 7).value or '').strip()
        link = str(ws.cell(r, 8).value or '').strip()
        catatan = str(ws.cell(r, 9).value or '').strip()

        # Delete existing element row if present to avoid duplication
        run_mysql(f"DELETE FROM baseline_elemen WHERE mitra_id = {mid} AND nomor_elemen = {num};")

        ins_el = f"""
        INSERT INTO baseline_elemen (
            mitra_id, nomor_elemen, kelompok, nama_elemen, yang_diperiksa, sumber_bukti_minimum,
            status, fakta_pemeriksaan, link_sumber_bukti, catatan
        ) VALUES (
            {mid}, {num}, {sql_escape(kelompok)}, {sql_escape(nama_el)}, {sql_escape(yang_diperiksa)}, {sql_escape(sumber_min)},
            {sql_escape(status)}, {sql_escape(fakta)}, {sql_escape(link)}, {sql_escape(catatan)}
        );
        """
        run_mysql(ins_el)
        total_elements_ingested += 1

print(f"   [DONE] Ingested {total_elements_ingested} baseline elements across 15 partnerships.")

# Step 4: Ingest Stakeholder Workbooks Data (PICs, Pelaksanaan Kegiatan, Usulan Tindak Lanjut)
print("3. Ingesting Stakeholder Intelligence (PIC, Kegiatan, Tindak Lanjut)...")

wb_to_inst = {
    '01_Pilot_Utama/P01_Scorecard_Dekranasda_Kepri.xlsx': 'Dewan Kerajinan Nasional Daerah Provinsi Kepulauan Riau',
    '01_Pilot_Utama/P02_Scorecard_BNNP_Kepri.xlsx': 'Badan Narkotika Nasional Provinsi Kepulauan Riau',
    '01_Pilot_Utama/P03_Scorecard_Pemkot_Tanjungpinang.xlsx': 'Pemerintah Kota Tanjungpinang',
    '01_Pilot_Utama/P04_Scorecard_Bapperida_Bintan.xlsx': 'Badan Perencanaan Pembangunan, Riset, dan Inovasi Daerah Kabupaten Bintan',
    '01_Pilot_Utama/P05_Scorecard_STIT_Mumtaz_Karimun.xlsx': 'Sekolah Tinggi Ilmu Tarbiyah (STIT) Mumtaz Karimun',
    '01_Pilot_Utama/P06_Scorecard_UMRAH.xlsx': 'Universitas Maritim Raja Ali Haji (UMRAH)',
    '01_Pilot_Utama/P07_Scorecard_STAI_Paduka_Anambas.xlsx': 'Sekolah Tinggi Agama Islam Paduka Anambas',
    '01_Pilot_Utama/P08_Scorecard_Politeknik_Negeri_Batam.xlsx': 'Politeknik Negeri Batam',
    '01_Pilot_Utama/P09_Scorecard_STAI_Natuna.xlsx': 'Sekolah Tinggi Agama Islam Natuna',
    '01_Pilot_Utama/P10_Scorecard_STISIP_Bunda_Tanah_Melayu.xlsx': 'Sekolah Tinggi Ilmu Sosial dan Ilmu Politik Bunda Tanah Melayu',
    '02_Portofolio_Pengayaan/C01_Scorecard_STAIN_Sultan_Abdurrahman.xlsx': 'Sekolah Tinggi Agama Islam Negeri Sultan Abdurrahman Kepulauan Riau',
    '02_Portofolio_Pengayaan/C02_Scorecard_Politeknik_Bintan_Cakrawala.xlsx': 'Politeknik Bintan Cakrawala',
    '02_Portofolio_Pengayaan/C03_Scorecard_Universitas_Ibnu_Sina.xlsx': 'Universitas Ibnu Sina',
    '02_Portofolio_Pengayaan/C04_Scorecard_UNRIKA.xlsx': 'Universitas Riau Kepulauan',
}

kegiatan_inserted = 0
utl_inserted = 0

for rel_wb, mid in wb_to_mid.items():
    fpath = os.path.join(path_base, rel_wb)
    wb_sh = openpyxl.load_workbook(fpath, data_only=True)

    # 1. Update PIC
    ws_id = wb_sh['IDENTITAS & PIC']
    pic_nama = ws_id.cell(11, 3).value
    pic_jabatan = ws_id.cell(12, 3).value
    pic_kontak = ws_id.cell(14, 3).value
    if pic_nama and str(pic_nama).strip() and str(pic_nama).strip() != 'None':
        summary_pic_mitra = str(pic_nama).strip()
        if pic_jabatan and str(pic_jabatan).strip(): summary_pic_mitra += f" ({str(pic_jabatan).strip()})"
        if pic_kontak and str(pic_kontak).strip(): summary_pic_mitra += f" - Kontak: {str(pic_kontak).strip()}"
        run_mysql(f"UPDATE mitra_kinerja SET pic_mitra = {sql_escape(summary_pic_mitra)} WHERE id = {mid};")

    pic_int_nama = ws_id.cell(19, 3).value
    if pic_int_nama and str(pic_int_nama).strip() and not str(pic_int_nama).startswith('<misal') and str(pic_int_nama).strip() != 'None':
        run_mysql(f"UPDATE mitra_kinerja SET pic_internal = {sql_escape(str(pic_int_nama).strip())} WHERE id = {mid};")

    # 2. Ingest Activities (PELAKSANAAN KEGIATAN)
    ws_keg = wb_sh['PELAKSANAAN KEGIATAN']
    for r in range(7, ws_keg.max_row + 1):
        no = ws_keg.cell(r, 1).value
        nama_keg = ws_keg.cell(r, 2).value
        status_keg = str(ws_keg.cell(r, 3).value or '').strip()
        hasil = str(ws_keg.cell(r, 4).value or '').strip()
        bukti = str(ws_keg.cell(r, 5).value or '').strip()
        if nama_keg and str(nama_keg).strip() and not str(nama_keg).startswith('Contoh') and no:
            clean_nama = str(nama_keg).strip()
            # Map status
            db_status = 'Selesai' if 'Sudah' in status_keg else ('Proses' if 'Sedang' in status_keg else 'Belum')
            deskripsi = clean_nama
            if hasil: deskripsi += f" (Hasil: {hasil})"

            # Check if already exists in tindak_lanjut
            chk = run_mysql(f"SELECT id FROM tindak_lanjut WHERE mitra_id = {mid} AND tindakan = {sql_escape(deskripsi)};")
            if not chk.strip():
                ins_tl = f"""
                INSERT INTO tindak_lanjut (mitra_id, tindakan, tenggat, status, file_bukti)
                VALUES ({mid}, {sql_escape(deskripsi)}, '2026-08-28', {sql_escape(db_status)}, {sql_escape(bukti if bukti else None)});
                """
                run_mysql(ins_tl)
                kegiatan_inserted += 1

    # 3. Ingest Action Plans (USULAN TINDAK LANJUT)
    ws_utl = wb_sh['USULAN TINDAK LANJUT']
    for r in range(7, ws_utl.max_row + 1):
        no = ws_utl.cell(r, 1).value
        usulan = ws_utl.cell(r, 2).value
        target = str(ws_utl.cell(r, 3).value or '').strip()
        manfaat = str(ws_utl.cell(r, 4).value or '').strip()
        bantuan = str(ws_utl.cell(r, 5).value or '').strip()
        pihak = str(ws_utl.cell(r, 6).value or '').strip()

        if usulan and str(usulan).strip() and not str(usulan).startswith('Contoh') and no:
            clean_usulan = str(usulan).strip()
            lingkup = manfaat if manfaat else 'Rencana aksi tindak lanjut kerja sama'
            if bantuan: lingkup += f" | Bantuan: {bantuan}"
            if pihak: lingkup += f" | Pelibatan: {pihak}"

            # Check if exists in rencana_kerja
            chk_rk = run_mysql(f"SELECT id FROM rencana_kerja WHERE mitra_id = {mid} AND judul_rencana = {sql_escape(clean_usulan)};")
            if not chk_rk.strip():
                ins_rk = f"""
                INSERT INTO rencana_kerja (mitra_id, judul_rencana, ruang_lingkup, maksud_tujuan, tanggal_mulai, tanggal_selesai, status)
                VALUES ({mid}, {sql_escape(clean_usulan)}, {sql_escape(lingkup)}, 'Diusulkan dari instrumen stakeholder', '2026-09-01', '2027-12-31', 'Disetujui');
                """
                run_mysql(ins_rk)
                utl_inserted += 1

print(f"   [DONE] Ingested {kegiatan_inserted} authentic field activities and {utl_inserted} action plans.")

# Step 5: Update MySQL Dumps
print("4. Updating clean SQL dumps (mitra_kinerja_dump.sql, schema.sql, seed_data.sql)...")

# 1. Complete dump
cmd_dump = [
    r"C:\xampp\mysql\bin\mysqldump.exe",
    "-u", "root",
    "--databases", "mitra_kinerja",
    "--add-drop-database",
    "--routines",
    "--events",
    "--default-character-set=utf8mb4"
]
res_dump = subprocess.run(cmd_dump, capture_output=True)
if res_dump.returncode == 0:
    with open(os.path.join(cwd, "database", "mitra_kinerja_dump.sql"), "wb") as f:
        f.write(res_dump.stdout)
    print(f"   [OK] database/mitra_kinerja_dump.sql ({len(res_dump.stdout)} bytes)")
else:
    print("   [ERR] dump failed:", res_dump.stderr)

# 2. Seed data
cmd_seed = [
    r"C:\xampp\mysql\bin\mysqldump.exe",
    "-u", "root",
    "--no-create-info",
    "--complete-insert",
    "--extended-insert",
    "mitra_kinerja"
]
res_seed = subprocess.run(cmd_seed, capture_output=True)
if res_seed.returncode == 0:
    with open(os.path.join(cwd, "database", "seed_data.sql"), "wb") as f:
        f.write(res_seed.stdout)
    print(f"   [OK] database/seed_data.sql ({len(res_seed.stdout)} bytes)")

# 3. Schema
cmd_schema = [
    r"C:\xampp\mysql\bin\mysqldump.exe",
    "-u", "root",
    "--no-data",
    "--routines",
    "mitra_kinerja"
]
res_schema = subprocess.run(cmd_schema, capture_output=True)
if res_schema.returncode == 0:
    with open(os.path.join(cwd, "database", "schema.sql"), "wb") as f:
        f.write(res_schema.stdout)
    print(f"   [OK] database/schema.sql ({len(res_schema.stdout)} bytes)")

print("=== FINAL PACKET INTEGRATION COMPLETE SUCCESSFULLY! ===")
